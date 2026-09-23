from collections import defaultdict
from pathlib import Path
import re

from openpyxl import load_workbook


OFFICIAL_MIN_YEAR = 2012
OFFICIAL_MAX_YEAR = 2024


def normalize_name(value):
    if value is None:
        return ""
    return " ".join(str(value).split())


def name_key(value):
    # "광주광역시 (구)"와 "광주광역시(구)"를 같은 이름으로 처리
    return re.sub(r"\s+", "", normalize_name(value))


def load_koroad_request_codes(xlsx_path: Path):
    workbook = load_workbook(
        xlsx_path,
        read_only=False,
        data_only=True,
    )

    # 실제 파일의 시트명은 serachYearCd로 오타가 나 있음
    year_sheet = workbook["serachYearCd 요청값"]
    sido_sheet = workbook["Sido 요청값"]
    gugun_sheet = workbook["Gugun 요청값"]

    # ---------------------------------------------------------
    # 1. 시도 코드
    # ---------------------------------------------------------
    sido_by_name = {}

    for row in range(2, sido_sheet.max_row + 1):
        name = sido_sheet.cell(row, 1).value
        code = sido_sheet.cell(row, 2).value

        if name is None or code is None:
            continue

        sido_by_name[name_key(name)] = f"{int(code):02d}"

    # ---------------------------------------------------------
    # 2. 자전거 연도와 예상 afos_id
    # ---------------------------------------------------------
    afos_id_by_year = {}
    current_category = ""

    for row in range(2, year_sheet.max_row + 1):
        category = year_sheet.cell(row, 1).value
        label = year_sheet.cell(row, 2).value
        dataset_id = year_sheet.cell(row, 3).value

        if category is not None:
            current_category = normalize_name(category)

        if current_category != "자전거 교통사고 다발지역":
            continue

        match = re.match(r"(\d{2})년", normalize_name(label))
        if not match:
            continue

        year = 2000 + int(match.group(1))

        if OFFICIAL_MIN_YEAR <= year <= OFFICIAL_MAX_YEAR:
            afos_id_by_year[year] = str(int(dataset_id))

    expected_years = list(
        range(OFFICIAL_MIN_YEAR, OFFICIAL_MAX_YEAR + 1)
    )

    if sorted(afos_id_by_year) != expected_years:
        raise ValueError(
            "자전거 제공연도 코드가 2012~2024와 일치하지 않습니다."
        )

    # ---------------------------------------------------------
    # 3. 시도 블록별 시군구 코드
    # ---------------------------------------------------------
    gugun_by_province = defaultdict(list)
    current_province = None

    for row in range(2, gugun_sheet.max_row + 1):
        province = gugun_sheet.cell(row, 1).value
        district = gugun_sheet.cell(row, 2).value
        gugun_code = gugun_sheet.cell(row, 3).value

        if province is not None:
            current_province = normalize_name(province)

        if district is None or gugun_code is None:
            continue

        gugun_by_province[current_province].append({
            "district_name": normalize_name(district),
            "guGun": f"{int(gugun_code):03d}",
        })

    # ---------------------------------------------------------
    # 4. 연도별 시도코드 결정
    # ---------------------------------------------------------
    def resolve_sido(year, province):
        province_key = name_key(province)

        if province_key == name_key("강원특별자치도"):
            source_name = (
                "강원도(구)"
                if year <= 2022
                else "강원특별자치도"
            )
            return sido_by_name[name_key(source_name)]

        if province_key == name_key("전북특별자치도"):
            source_name = (
                "전라북도(구)"
                if year <= 2022
                else "전북특별자치도"
            )
            return sido_by_name[name_key(source_name)]

        return sido_by_name[province_key]

    # ---------------------------------------------------------
    # 5. 전국·전연도 요청 범위 생성
    # ---------------------------------------------------------
    request_scopes = []

    for year in expected_years:
        for province, districts in gugun_by_province.items():
            si_do = resolve_sido(year, province)

            # 2012~2024 자전거 API 수집 대상에서 제외
            if si_do == "12":
                continue

            for district in districts:
                request_scopes.append({
                    "searchYearCd": str(year),
                    "siDo": si_do,
                    "guGun": district["guGun"],
                    "province_name": province,
                    "district_name": district["district_name"],
                    "expected_afos_id": afos_id_by_year[year],
                })

    if len(request_scopes) != 3510:
        raise ValueError(
            f"예상 요청범위 3,510개와 다릅니다: "
            f"{len(request_scopes):,}개"
        )

    return request_scopes, afos_id_by_year

if __name__ == "__main__":
    code_list_path = (
        Path(__file__).resolve().parent
        / "AccidentHazard_CodeList.xlsx"
    )

    if not code_list_path.exists():
        raise FileNotFoundError(
            f"코드표 파일이 없습니다: {code_list_path}"
        )

    request_scopes, afos_id_by_year = (
        load_koroad_request_codes(code_list_path)
    )

    print("=" * 65)
    print("KoROAD 자전거 교통사고 전국 요청범위 생성")
    print("=" * 65)
    print(f"코드표       : {code_list_path}")
    print(f"수집연도     : {min(afos_id_by_year)}~{max(afos_id_by_year)}")
    print(f"요청범위     : {len(request_scopes):,}개")
    print()

    print("[연도별 afos_id 검증값]")
    for year, afos_id in sorted(afos_id_by_year.items()):
        print(f"{year} → {afos_id}")

    print()
    print("[처음 5개 요청범위]")

    for index, scope in enumerate(request_scopes[:5], start=1):
        print(
            f"{index}. "
            f"{scope['searchYearCd']} / "
            f"{scope['province_name']}({scope['siDo']}) / "
            f"{scope['district_name']}({scope['guGun']})"
        )

    print()
    print("[마지막 요청범위]")

    last_scope = request_scopes[-1]

    print(
        f"{last_scope['searchYearCd']} / "
        f"{last_scope['province_name']}({last_scope['siDo']}) / "
        f"{last_scope['district_name']}({last_scope['guGun']})"
    )

    print()
    print("=" * 65)
    print("전국·전연도 요청범위 생성 성공")
    print("=" * 65)
